from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path('reports/RiskPilot_AI_Competition_Report.docx')
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.75)
sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
sec.header_distance=Inches(.32);sec.footer_distance=Inches(.32)
sec.different_first_page_header_footer=False
styles=doc.styles
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
    st=styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st.font.underline=False
    st.paragraph_format.widow_control=True
normal=styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',27),('Heading 1',19),('Heading 2',13),('Heading 3',11)]:
    styles[name].font.size=Pt(size)
    styles[name].paragraph_format.space_before=Pt(13 if name!='Title' else 0)
    styles[name].paragraph_format.space_after=Pt(8)
    styles[name].paragraph_format.keep_with_next=True
styles['Title'].font.bold=False
styles['Subtitle'].font.size=Pt(12)
styles['Subtitle'].font.italic=False
for style_name in ('Title','Subtitle','Title Char','Subtitle Char'):
    if style_name in styles:
        element=styles[style_name]._element
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        for spacing in list(element.iter(qn('w:spacing'))):
            if spacing.getparent().tag==qn('w:rPr'):spacing.getparent().remove(spacing)
header=sec.header.paragraphs[0]
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');header._p.append(f)
footer=sec.footer.paragraphs[0]

doc.core_properties.title='RiskPilot AI SOC Triage Technical Report'
doc.core_properties.author='Wong Li En; Tian Xinyu; Wong Chin Wen'
doc.core_properties.subject='HackAI 2026 Agentic Cyber Defence technical report'


def p(text,style=None):
    return doc.add_paragraph(text,style)
def h(text,level=1):
    return doc.add_heading(text,level)
def page():
    doc.add_page_break()
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for idx,values in enumerate([headers]+rows):
        row=t.rows[0] if idx==0 else t.add_row()
        if idx==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for j,(cell,text) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for edge,value in [('top','95'),('bottom','95'),('left','110'),('right','110')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');m.append(e)
            tcPr.append(m)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8E8E8' if idx==0 else ('F4F4F4' if idx%2==0 else 'FFFFFF'));tcPr.append(shade)
            para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
            para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(str(text))>24 else WD_ALIGN_PARAGRAPH.CENTER
            run=para.add_run(str(text));run.font.size=Pt(10);run.font.bold=idx==0
    p('').paragraph_format.space_after=Pt(1)
    return t


cover_title=p('RiskPilot AI Agentic Cyber Defence Report', 'Title')
cover_title.paragraph_format.space_before=Pt(72)
cover_title.paragraph_format.space_after=Pt(22)
cover_title.paragraph_format.line_spacing=2
cover_title.alignment=WD_ALIGN_PARAGRAPH.CENTER
for run in cover_title.runs:
    run.font.size=Pt(11)
    run.font.bold=True
for text in ('Wong Li En, Tian Xinyu, and Wong Chin Wen', "Taylor's University", 'HackAI 2026', 'Cyber Defence Track', 'October 24, 2026'):
    para=p(text)
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing=2
    para.paragraph_format.space_after=Pt(0)
page()
p('RiskPilot AI organises historical network alerts into evidence-linked cases and uses a bounded AI agent to investigate them. It produces an explainable assessment and a proposed investigation plan for an analyst to review. The demonstrated scope covers PortScan, DoS and suspected BruteForce activity.')
h('01 Problem and user')
h('The triage problem',2)
p('A network alert identifies behaviour that matches a detection rule. An analyst still needs to establish what happened, which evidence supports the signal and what to investigate next. Reviewing isolated alerts makes it difficult to follow activity over time or distinguish an observed traffic pattern from an inferred attack.')
p('The same source may generate alerts across multiple minutes. Treating every alert as a separate investigation fragments the evidence, while treating a rule match as proof of compromise can produce an unsupported conclusion. Effective triage needs both correlation and a clear way to express uncertainty.')
h('Intended users',2)
p('The prototype supports Tier 1 and Tier 2 security operations centre analysts. A first reviewer can understand the trigger, inspect the timeline and use the AI findings to identify the next useful check. A more experienced analyst can review the cited records, accept or reject the assessment and decide whether further investigation is justified.')
h('Project objective',2)
p('We built a case-centred investigation workflow that connects rule alerts, traffic evidence, AI reasoning and analyst decisions. Its defensive output is a prioritised review queue and an evidence-linked investigation plan. The prototype replays historical flow data; operational response remains under human control.')

page()
h('02 Inputs and evidence')
h('Data and scope',2)
p('Development uses labelled network-flow data from CIC-IDS2017 (Canadian Institute for Cybersecurity, n.d.; Sharafaldin et al., 2018). The selected files cover Monday, Tuesday, Wednesday and Friday afternoon PortScan traffic. The application displays three detection categories: PortScan, DoS and BruteForce. Benign is used only as an offline evaluation label; it is not a detection rule or an automatic safe verdict.')
p('The demonstration contains 9 grouped cases, 158 alerts and 948 sampled flow records. Its 1,106 timeline events link to an alert or retained sample. Each sample retains its source-file and row reference so the analyst can trace it to the underlying data.')
h('Detection basis',2)
p('Rules operate on a source file, ordered source and destination IP pair, and calendar minute. One CSV row counts as one flow record; forward and backward packet counts are summed. The following thresholds are prototype heuristics.')
table(['Category','Rule signal within one minute'],[
['PortScan','At least 100 flows and 100 destination ports, plus either source ports at least twice destination ports, or at least 250 destination ports and 2 source ports.'],
['DoS','At least 1,000 flows and 10,000 packets, at most 10 destination ports and at least 10 source ports.'],
['BruteForce','At least 50 connections from one source to one target FTP port 21 or SSH port 22.']],[1.1,5.45])
p('A rule signal warrants investigation. Repeated connections to a login port do not establish failed passwords, and traffic volume does not establish service disruption. An unmatched record is not automatically classified as Benign.')
h('Case grouping and evidence access',2)
p('Alerts are grouped when their source file, rule type, ordered address pair and time proximity match, with at most five minutes between alert-minute starts. BruteForce grouping also requires the same target login port. Separate sources or destinations remain separate cases, although related cases can be linked.')
p('The interface provides minute-level drill-down, sampled evidence and paginated access to locally indexed original flow records. AI tools receive case-scoped aggregates and bounded samples. Dataset labels and analyst-only annotations are withheld from the model. Source timestamps are shown as dataset time rather than assuming an unverified timezone.')

page()
h('03 Agent workflow')
h('Investigation process',2)
p('The analyst opens a case and starts an AI investigation, or enables sequential investigation of pending cases. The AI first reads the case details, then selects evidence queries according to what it needs to verify. The backend executes only allowlisted read-only tools and returns the results to the model. The agent can request further evidence before producing its proposal.')
table(['Read-only tool','Purpose'],[
['Case details','Read the addresses, complete grouped time range, flow total and evidence scope.'],
['Related alerts','Inspect the rule measurements and linked alert IDs, with continuation for longer lists.'],
['Network flows','Inspect sampled communication records, including ports, packet counts and evidence IDs.'],
['Incident statistics','Read case aggregates and statistics explicitly derived from the retained samples.']],[1.65,4.9])
p('The investigation is limited to six model steps and twelve tool calls. Tools are scoped to the active case. Unknown tools, invalid arguments and requests for another case are refused. Bounded tool responses preserve valid structured data and identify omitted records. Provider failures produce an unavailable state rather than a substituted mock conclusion.')
h('Assessment and evidence checks',2)
p('The agent proposes Likely Attack, Likely Benign or Needs Review, together with its reasoning, model-reported confidence, cited records, findings, limitations and up to two next steps. Observation and inference are presented separately. Confidence is not a calibrated probability of attack.')
p('The backend checks the case identity, allowed output fields, evidence ownership, retrieved citations, supported numeric claims and specified time-window conflict claims. New results carry an evidence fingerprint so changed case evidence invalidates the cached proposal. Rejected proposals are withheld from confirmation and recommendation approval. These checks constrain output; they do not prove every interpretation correct.')
h('Timeline and ATT&CK context',2)
p('The timeline orders the case evidence and opens the relevant record from a finding or recommendation. PortScan evidence may support a tentative T1046 Network Service Discovery mapping (MITRE, 2026b). DoS and BruteForce remain Unmapped when the flow data cannot establish service impact or credential guessing, as described by T1498 (MITRE, 2025) and T1110 (MITRE, 2026a). Mapping hints remain distinguishable from confirmed attacks.')

page()
h('04 Defensive action')
p('RiskPilot AI helps an analyst decide where to focus and what evidence to examine next. The case view combines the rule signal, timeline, AI findings and proposed investigation steps. Each suggested step includes its rationale and links to supporting evidence.')
p('The queue uses an explainable review-priority heuristic based on traffic rate, related alert count, source diversity and destination-port spread. High, Medium and Low indicate review order. They are not calibrated severity levels or probabilities. Analysts can inspect the contributing reasons and record a justified manual adjustment in the audit trail.')
p('Permitted proposals are to review network flows, check related alerts, review the incident timeline or escalate to an analyst. The agent has no tool for blocking an address, isolating a host, modifying evidence or closing a case. Its authority is restricted to investigation and proposed next steps, consistent with limiting excessive agency (OWASP Foundation, n.d.).')
h('05 Human approval')
p('The system separates the AI proposal from the analyst decision. Viewing evidence is read-only; confirming a conclusion, approving a step and resolving a case are distinct operations.')
table(['Analyst decision','Meaning and recorded effect'],[
['Confirm or reject assessment','Records the analyst judgment about the AI interpretation. Rejection requires a reason; confirming Likely Benign on a protected attack category also requires a written reason.'],
['Approve or reject a step','Records selected investigation-plan decisions. Approved steps enter the pending work list and create a case-note draft. Approval does not execute the investigation.'],
['Save case record','Preserves the investigation summary as an audit snapshot. Optional unreviewed steps remain visible as such. Saving does not resolve the case.'],
['Resolve case','Requires a saved record, an explicit analyst confirmation and a written reason. Resolution removes the case from the active review queue.']],[1.75,4.8])
p('Approved work is completed by the analyst and marked done separately. The case-note draft remains a draft. The printable case record includes case facts, detection evidence, the AI assessment, analyst review, investigation-plan decisions and activity history.')
p('Write operations use an authenticated session and CSRF protection, and the reviewer must match the signed-in user. AI activity, validation outcomes, analyst decisions, drafts and saved records enter a shared audit trail. A hash chain and database update/delete restrictions support append-only auditing through normal application access.')

page()
h('06 Result and limitations')
h('Working prototype',2)
p('The application demonstrates grouped case investigation across 9 cases, with all 1,106 exported timeline events resolving to case evidence. The software suite passed 280 automated tests, including evidence validation, case boundaries, reviewer identity, safe resolution and audit behaviour. A separate set of 21 synthetic safety scenarios passed. These are functional checks rather than measures of detection accuracy or analyst effectiveness.')
h('Detection results',2)
p('Development scoring evaluated 356,103 minute groups, excluding 11 out-of-scope groups. Across the target attack types, the detector recorded 156 true positives, 2 false positives and 73 false negatives: binary precision 98.73% and recall 68.12%. The Benign false-positive rate was 2 of 355,863 groups, approximately 0.000562%. These data informed development and are not an independent test.')
table(['Category','Development precision','Development recall','External recall'],[
['PortScan','91.67%','40.74%','0.118%'],
['DoS','95.83%','30.67%','41.176%'],
['BruteForce','100.00%','96.06%','0.359%']],[1.1,1.8,1.8,1.85])
p('The independent-source detector evaluation used CIDDS-001 OpenStack internal weeks 1 to 4 (Hochschule Coburg, n.d.). It scored 5,942,316 minute groups after excluding 4,260 ambiguous or out-of-scope groups. The thresholds were fixed before scoring. The low external recall shows that the current rules do not generalise reliably, especially for PortScan and BruteForce. Zero observed external false positives must be read alongside the small number of positive predictions.')
h('Real AI investigation example',2)
p('A Codex investigation examined a case containing 61 BruteForce-rule alerts between 09:19 and 10:19 on 4 July 2017. It completed in 54.89 seconds using four read-only tool calls and four cited records. The agent returned Needs Review with medium confidence: repeated connections warranted investigation, but the evidence could establish neither credential guessing nor a legitimate purpose.')
p('The agent proposed reviewing the linked flow patterns and escalating for authentication-log and authorised-activity checks. Its output passed evidence validation. An isolated workflow test recorded a simulated analyst confirmation, one approved step, a DRAFT case note and a saved case record. The resulting ten-event audit chain verified successfully. This single case demonstrates the workflow; it does not establish AI classification accuracy or human decision quality.')

page()
h('Limitations',2)
p('Detection coverage is the main limitation. Attacks missed by the initial rules do not reach the normal AI investigation queue. Fixed minute boundaries can split activity, simple count thresholds can miss slow behaviour, and first-match rule priority can mask overlapping signals. Address-and-time grouping is a practical correlation method, not proof of a common campaign.')
p('Flow samples lack authentication outcomes, payload content, asset importance, business context and service-health evidence. These gaps limit conclusions about BruteForce, DoS impact and intent. ATT&CK hints remain tentative, and incomplete evidence should result in further review rather than a forced attack or benign decision.')
p('AI confidence is model-reported and uncalibrated. Deterministic checks verify selected evidence and claim constraints, not all reasoning or causal assertions. The demonstrated Codex configuration uses a default model alias rather than a pinned underlying version. Broader evaluation needs separate external cases, expert review and measurement of both errors and abstention.')
p('The prototype has not measured investigation-time savings, workload reduction or improved analyst outcomes. Approval clicks and written reasons establish accountability, not the correctness of the security judgment. The locally verifiable audit chain is also not tamper-proof against an administrator able to replace the application and storage.')
p('Future work should strengthen evidence coverage, evaluate detection on disjoint calibration and test data, and measure usability with analysts. The current contribution is a working, constrained investigation workflow with traceable evidence and explicit human decisions.')
page()
ref_heading=h('References',2)
ref_heading.alignment=WD_ALIGN_PARAGRAPH.CENTER
ref_heading.paragraph_format.space_before=Pt(0)
ref_heading.paragraph_format.space_after=Pt(0)
ref_heading.paragraph_format.line_spacing=2
references=[
('Canadian Institute for Cybersecurity. (n.d.). ', 'Intrusion detection evaluation dataset (CIC-IDS2017)', '. University of New Brunswick. https://www.unb.ca/cic/datasets/ids-2017.html'),
('Hochschule Coburg. (n.d.). ', 'CIDDS: Coburg intrusion detection data sets', '. https://www.hs-coburg.de/forschen/cidds-coburg-intrusion-detection-data-sets/'),
('MITRE. (2025, October 24). ', 'Network denial of service (T1498)', '. MITRE ATT&CK. https://attack.mitre.org/techniques/T1498/'),
('MITRE. (2026a, May 12). ', 'Brute force (T1110)', '. MITRE ATT&CK. https://attack.mitre.org/techniques/T1110/'),
('MITRE. (2026b, May 12). ', 'Network service discovery (T1046)', '. MITRE ATT&CK. https://attack.mitre.org/techniques/T1046/'),
('OWASP Foundation. (n.d.). ', 'LLM06:2025 excessive agency', '. OWASP Gen AI Security Project. https://genai.owasp.org/llmrisk/llm062025-excessive-agency/'),
('Sharafaldin, I., Lashkari, A. H., & Ghorbani, A. A. (2018). Toward generating a new intrusion detection dataset and intrusion traffic characterization. In P. Mori, S. Furnell, & O. Camp (Eds.), ', 'Proceedings of the 4th International Conference on Information Systems Security and Privacy', ' (Vol. 1, pp. 108\u2013116). SciTePress. https://doi.org/10.5220/0006639801080116')
]
for prefix,title,suffix in references:
    para=p('')
    para.paragraph_format.left_indent=Inches(.5)
    para.paragraph_format.first_line_indent=Inches(-.5)
    para.paragraph_format.space_after=Pt(0)
    para.paragraph_format.line_spacing=2
    para.add_run(prefix)
    para.add_run(title).italic=True
    before_url,url=suffix.split('https://',1)
    para.add_run(before_url.rstrip())
    para.add_run().add_break()
    para.add_run('https://'+url)

# Explicit black text throughout, independent of Word theme settings.
for element in doc.element.iter(qn('w:r')):
    from docx.text.run import Run
    Run(element, None).font.color.rgb=RGBColor(0,0,0)
for part in (sec.header,sec.footer,sec.first_page_header,sec.first_page_footer):
    for para in part.paragraphs:
        for run in para.runs:run.font.color.rgb=RGBColor(0,0,0)
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
